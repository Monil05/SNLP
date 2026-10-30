"""Serve CareerMatch and proxy /api/* requests to the local career_nlp API."""
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen
import json
import os

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
ALLOWED_DOCUMENT_SUFFIXES = {".pdf", ".doc", ".docx"}

FRONTEND_DIR = Path(__file__).resolve().parent
HOST = "127.0.0.1"
PORT = int(os.environ.get("FRONTEND_PORT", "5500"))
BACKEND_URL = "http://127.0.0.1:8000"
GROUP7_API_URL = os.environ.get("GROUP7_API_URL", "http://127.0.0.1:8080/api").rstrip("/")


class FrontendHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(FRONTEND_DIR), **kwargs)

    def do_GET(self):
        path = urlsplit(self.path).path
        if path.startswith("/api/"):
            self._proxy_api()
        else:
            super().do_GET()

    def do_POST(self):
        if urlsplit(self.path).path == "/extract-document":
            self._extract_document()
        elif urlsplit(self.path).path.startswith("/api/"):
            self._proxy_api()
        else:
            self.send_error(405, "POST is only supported for document extraction and /api/ requests")

    def _extract_document(self):
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0:
            self._send_json(400, {"detail": "Choose a PDF or Word document to upload."})
            return
        if length > MAX_UPLOAD_BYTES + 1024 * 1024:
            self._send_json(413, {"detail": "Files must be 5 MB or smaller."})
            return

        body = self.rfile.read(length)
        content_type = self.headers.get("Content-Type", "")
        message = BytesParser(policy=policy.default).parsebytes(
            f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("ascii") + body
        )
        upload = next((part for part in message.iter_parts() if part.get_param("name", header="content-disposition") == "file"), None)
        if upload is None or not upload.get_filename():
            self._send_json(400, {"detail": "The uploaded document was not received."})
            return

        filename = Path(upload.get_filename()).name
        suffix = Path(filename).suffix.lower()
        data = upload.get_payload(decode=True) or b""
        if suffix not in ALLOWED_DOCUMENT_SUFFIXES:
            self._send_json(415, {"detail": "Upload a PDF (.pdf) or Word document (.doc or .docx)."})
            return
        if len(data) > MAX_UPLOAD_BYTES:
            self._send_json(413, {"detail": "Files must be 5 MB or smaller."})
            return
        if not data:
            self._send_json(400, {"detail": "The selected file is empty."})
            return

        try:
            extracted = self._read_document(data, suffix)
            extracted = "\n".join(line.strip() for line in extracted.splitlines() if line.strip()).strip()
            if not extracted:
                raise ValueError("No readable text was found. If this is a scanned PDF, run OCR or paste its text manually.")
            self._send_json(200, {"filename": filename, "text": extracted})
        except (ValueError, RuntimeError) as error:
            self._send_json(422, {"detail": str(error)})

    @staticmethod
    def _read_document(data, suffix):
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
                return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
            except ImportError as error:
                raise RuntimeError("PDF extraction requires pypdf. Install it with: pip install pypdf") from error
            except Exception as error:
                raise ValueError(f"Could not read this PDF: {error}") from error

        if suffix == ".docx":
            try:
                from docx import Document
                document = Document(BytesIO(data))
                paragraphs = [paragraph.text for paragraph in document.paragraphs]
                paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                return "\n".join(paragraphs)
            except ImportError as error:
                raise RuntimeError("Word extraction requires python-docx. Install it with: pip install python-docx") from error
            except Exception as error:
                raise ValueError(f"Could not read this Word document: {error}") from error

        with tempfile.TemporaryDirectory(prefix="careermatch-doc-") as temp_dir:
            source = Path(temp_dir) / "upload.doc"
            source.write_bytes(data)
            antiword = shutil.which("antiword")
            try:
                if antiword:
                    result = subprocess.run([antiword, str(source)], capture_output=True, text=True, timeout=30)
                    if result.returncode == 0 and result.stdout.strip():
                        return result.stdout
                office = shutil.which("soffice") or shutil.which("libreoffice")
                if office:
                    result = subprocess.run(
                        [office, "--headless", "--convert-to", "txt:Text", "--outdir", temp_dir, str(source)],
                        capture_output=True, text=True, timeout=45
                    )
                    converted = Path(temp_dir) / "upload.txt"
                    if result.returncode == 0 and converted.exists():
                        return converted.read_text(encoding="utf-8", errors="replace")
            except subprocess.TimeoutExpired as error:
                raise RuntimeError("The .doc conversion timed out. Save the document as .docx and upload it again.") from error
        raise RuntimeError("Legacy .doc extraction needs LibreOffice or antiword installed. Save the file as .docx, or install LibreOffice and restart the frontend server.")

    def _proxy_api(self):
        parsed = urlsplit(self.path)
        if parsed.path.startswith("/api/group7/"):
            api_path = parsed.path[len("/api/group7"):]
            target_base = GROUP7_API_URL
        else:
            api_path = parsed.path[len("/api"):]
            target_base = BACKEND_URL

        if not api_path.startswith("/v1/"):
            self._send_json(404, {"detail": "Unknown API path."})
            return

        target_url = f"{target_base}{api_path}"
        if parsed.query:
            target_url += f"?{parsed.query}"

        body = None
        length = int(self.headers.get("Content-Length", "0"))
        if length:
            body = self.rfile.read(length)

        headers = {"Accept": self.headers.get("Accept", "application/json")}
        content_type = self.headers.get("Content-Type")
        if content_type:
            headers["Content-Type"] = content_type
        api_key = self.headers.get("X-API-Key")
        if api_key:
            headers["X-API-Key"] = api_key

        request = Request(target_url, data=body, headers=headers, method=self.command)
        try:
            with urlopen(request, timeout=180) as response:
                payload = response.read()
                status = response.status
                response_type = response.headers.get("Content-Type", "application/json")
        except HTTPError as error:
            payload = error.read()
            status = error.code
            response_type = error.headers.get("Content-Type", "application/json")
        except (URLError, TimeoutError, OSError) as error:
            reason = getattr(error, "reason", error)
            self._send_json(502, {
                "detail": (
                    f"Could not connect to the service at {target_base}. "
                    "Check that the relevant API is running in another terminal, then try again. "
                    f"({reason})"
                )
            })
            return

        self.send_response(status)
        self.send_header("Content-Type", response_type)
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(payload)

    def _send_json(self, status, value):
        payload = json.dumps(value).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)


if __name__ == "__main__":
    os.chdir(FRONTEND_DIR)
    server = ThreadingHTTPServer((HOST, PORT), FrontendHandler)
    print(f"CareerMatch frontend: http://{HOST}:{PORT}")
    print(f"Proxying /api/* requests to {BACKEND_URL}")
    print(f"Proxying /api/group7/* requests to {GROUP7_API_URL}")
    print("Press Ctrl+C to stop.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nFrontend server stopped.")
    finally:
        server.server_close()
