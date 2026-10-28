"""Tests for the Week 6 reusable asset (career_nlp).

Offline regression tests for the Week 6 asset. Most use a small synthetic corpus;
one end-to-end sample test also checks ranking against the bundled full job corpus.
Run:  pytest -q test/test_career_nlp.py
"""
import io
import json
import sys
import types
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(SRC))

from career_nlp import (  # noqa: E402
    AssetConfig, CareerNLP, CVProcessingError, InputValidationError, ModelLoadError,
)
from career_nlp.embedders import SentenceTransformerEmbedder, find_local_snapshot  # noqa: E402
from career_nlp.extraction import extract_academic_records, extract_degrees, extract_phones, skill_confidence  # noqa: E402
from career_nlp.text import clean_job_title, parse_skills  # noqa: E402

# --------------------------------------------------------------------------
# fixtures
# --------------------------------------------------------------------------
_TEMPLATES = {
    "INFORMATION-TECHNOLOGY": (["python", "sql", "machine learning", "data analysis", "power bi", "aws"],
                               ["software engineer", "data analyst", "systems administrator"]),
    "FINANCE": (["financial reporting", "budgeting", "forecasting", "excel", "accounting", "sap"],
                ["finance analyst", "accountant", "financial controller"]),
    "HR": (["recruitment", "onboarding", "payroll", "employee relations", "talent acquisition", "hris"],
           ["hr generalist", "recruiter", "hr manager"]),
    "SALES": (["lead generation", "negotiation", "crm", "cold calling", "account management", "forecasting"],
              ["sales representative", "account executive", "sales manager"]),
}


def _make_jobs(per_category: int = 12) -> pd.DataFrame:
    rows = []
    for cat, (skills, titles) in _TEMPLATES.items():
        for i in range(per_category):
            chosen = [skills[(i + k) % len(skills)] for k in range(4)]
            title = titles[i % len(titles)]
            if i % 5 == 0:
                title += " II"
            rows.append({
                "job_id": f"job_{cat[:3]}_{i:03d}", "category": cat, "job_title": title,
                "job_description": f"We are hiring a {title}. Responsibilities involve {', '.join(chosen)} "
                                   f"and reporting to the {cat.lower().replace('-', ' ')} team. Requires {chosen[0]}.",
                "job_skill_set": str(chosen),
            })
    return pd.DataFrame(rows)


class HashEmbedder:
    """Deterministic bag-of-words hashing embedder (no dependencies)."""
    name, dimensions = "hash-64", 64

    def encode(self, texts):
        out = np.zeros((len(texts), self.dimensions), dtype=np.float32)
        for i, t in enumerate(texts):
            for w in t.lower().split():
                out[i, sum(map(ord, w)) % self.dimensions] += 1.0
        norms = np.linalg.norm(out, axis=1, keepdims=True); norms[norms == 0] = 1
        return out / norms


@pytest.fixture(scope="module")
def jobs():
    return _make_jobs()


@pytest.fixture(scope="module")
def nlp(jobs):
    return CareerNLP(AssetConfig(embedding_backend="lsa", lsa_dimensions=16), jobs=jobs).warmup()


IT_STUDENT = "B.Tech in Computer Science. Skills: python, sql, machine learning, data analysis. 2 years experience."
IT_JD = "Looking for a data analyst skilled in python, sql, power bi and aws."

# --------------------------------------------------------------------------
# structured output + confidence
# --------------------------------------------------------------------------
def test_lazy_init(jobs):
    n = CareerNLP(AssetConfig(embedding_backend="lsa", lsa_dimensions=16), jobs=jobs)
    assert not n.is_ready
    n.extract("python")
    assert n.is_ready


def test_match_structure_and_default_hides_internals(nlp):
    r = nlp.match(IT_STUDENT, IT_JD)
    d = json.loads(r.model_dump_json(exclude_none=True))
    assert 0 <= d["match_score"] <= 100
    assert {"python", "sql"} <= set(d["matched_skills"])
    assert {"power bi", "aws"} <= set(d["missing_weak_skills"])
    assert set(d["recommended_learning"]) == set(d["missing_weak_skills"])
    assert d["confidence"]["level"] in {"low", "medium", "high"} and d["confidence"]["calibrated"] is False
    assert "score_breakdown" not in d and "student_profile" not in d
    assert "factors" not in d["confidence"]
    assert d["meta"]["schema_version"] == "1.0" and d["meta"]["processing_ms"] >= 0


def test_match_details_opt_in(nlp):
    r = nlp.match(IT_STUDENT, IT_JD, include_details=True)
    b = r.score_breakdown
    assert (b.skill_weight, b.semantic_weight, b.category_weight) == (0.63, 0.27, 0.10)
    assert r.match_score == pytest.approx(
        0.63 * b.skill_match_percent
        + 0.27 * b.semantic_similarity_percent
        + 0.10 * b.category_compatibility_percent,
        abs=0.02,
    )
    assert r.confidence.factors and r.student_profile.skills


def test_match_score_formula_and_extremes(nlp):
    full = nlp.match("python sql power bi aws machine learning", "python sql power bi aws", include_details=True)
    assert full.score_breakdown.skill_match_percent == 100.0
    none = nlp.match("recruitment payroll onboarding", "python sql power bi aws", include_details=True)
    assert none.score_breakdown.skill_match_percent == 0.0
    assert full.match_score > none.match_score


def test_match_warns_when_jd_has_no_known_skills(nlp):
    r = nlp.match(IT_STUDENT, "We want a friendly person who enjoys cake.")
    assert any("No known skills were recognised in the job description" in w for w in r.meta.warnings)
    assert r.confidence.level == "low"


def test_learning_order_is_by_dataset_demand(nlp):
    r = nlp.match("recruitment", "python sql power bi aws machine learning data analysis")
    st = nlp._state
    demand = [st.skill_demand[s] for s in r.recommended_learning]
    assert demand == sorted(demand, reverse=True)


def test_recommend_structure(nlp):
    r = nlp.recommend(IT_STUDENT, top_k=3, include_details=True)
    assert r.recommended_job.rank == 1 and r.recommended_job.category == "INFORMATION-TECHNOLOGY"
    assert len(r.alternatives) == 2
    titles = [r.recommended_job.job_title] + [a.job_title for a in r.alternatives]
    assert len(set(t.lower() for t in titles)) == 3           # alternatives are title-distinct
    assert all(not t.endswith(" II") for t in titles)          # level suffix cleaned
    assert set(r.confidence.factors) == {"category_agreement_top5", "top_job_skill_coverage",
                                         "profile_skill_evidence", "score_margin_to_second"}
    assert r.student_profile is not None


def test_recommend_applies_subject_penalty_to_stored_job(jobs):
    custom_jobs = jobs.copy()
    custom_jobs.loc[0, "job_title"] = "Physics-focused data analyst"
    custom_jobs.loc[0, "job_description"] = (
        "Data analyst requiring strong expertise in Physics. "
        "Responsibilities include python, sql, machine learning and data analysis."
    )
    n = CareerNLP(AssetConfig(embedding_backend="lsa", lsa_dimensions=16), jobs=custom_jobs).warmup()
    result = n.recommend(
        "B.Tech Data Science. Skills: python, sql, machine learning, data analysis. Physics: 35%.",
        top_k=10,
        include_details=True,
    )
    all_jobs = [result.recommended_job, *result.alternatives]
    physics_job = next(job for job in all_jobs if job.job_id == str(custom_jobs.iloc[0]["job_id"]))
    # A Physics-heavy title plus a requirement is not enough to prove Physics
    # is the role's core domain when the responsibilities are general data work.
    assert physics_job.academic_penalty_percent == 40.0
    assert physics_job.match_score == pytest.approx(max(0.0, physics_job.base_match_score - 40.0))
    assert physics_job.academic_adjustments[0].domain_relevance_percent < 75.0


def test_recommend_no_skills_warns_and_is_low_confidence(nlp):
    r = nlp.recommend("I like long walks and good coffee")
    assert r.confidence.level == "low"
    assert any("No known skills" in w for w in r.meta.warnings)


def test_recommend_category_split_caps_confidence():
    from career_nlp import confidence as c
    split = c.recommend_confidence(80, 40, 100, category_agreement=0.4, n_student_skills=9)
    solid = c.recommend_confidence(80, 40, 100, category_agreement=1.0, n_student_skills=9)
    assert split.level == "low" and split.score < c.MEDIUM
    assert solid.level == "high"


def test_classify_confidence(nlp):
    r = nlp.classify("finance analyst: budgeting forecasting accounting excel sap financial reporting")
    assert r.predicted_category == "FINANCE"
    assert r.confidence.score == r.top_categories[0].probability
    assert r.confidence.method == "classifier_probability"
    assert sum(c.probability for c in nlp.classify("python sql", top_k=5).top_categories) == pytest.approx(1.0, abs=1e-3)


def test_analyze_dispatches(nlp):
    assert nlp.analyze(IT_STUDENT).flow == "jd_not_provided"
    assert nlp.analyze(IT_STUDENT, "  ").flow == "jd_not_provided"
    a = nlp.analyze(IT_STUDENT, IT_JD)
    assert a.flow == "jd_provided" and a.match and a.recommendation is None


def test_retrieve_embed_similarity_summarize(nlp):
    r = nlp.retrieve("python sql machine learning", top_k=3)
    assert [x.rank for x in r.results] == [1, 2, 3]
    assert r.results[0].similarity >= r.results[1].similarity >= r.results[2].similarity
    assert "Job " in r.context
    e = nlp.embed(["a python developer", "an hr recruiter"])
    assert len(e.vectors) == 2 and e.dimensions == len(e.vectors[0])
    s = nlp.similarity("python sql", "python sql")
    assert s.similarity == pytest.approx(1.0, abs=1e-3) and s.similarity_percent == pytest.approx(100, abs=0.1)
    text = "Python is used for data. SQL queries tables. Dashboards use power bi. Cats nap. Lunch is nice."
    sm = nlp.summarize(text, max_sentences=2)
    assert sm.sentences_total == 5 and sm.sentences_used == 2
    assert [x.index for x in sm.selected] == sorted(x.index for x in sm.selected)


def test_week3_compat_shape(nlp):
    d = CareerNLP.to_week3_dict(nlp.match(IT_STUDENT, IT_JD))
    assert set(d) == {"match_score", "matched_skills", "missing_weak_skills", "recommended_learning"}
    assert isinstance(d["matched_skills"], str) and "; " in d["matched_skills"]
    r = CareerNLP.to_week3_dict(nlp.recommend(IT_STUDENT))
    assert set(r) == {"recommended_job", "match_score", "matched_skills", "missing_weak_skills", "recommended_learning"}


# --------------------------------------------------------------------------
# extraction
# --------------------------------------------------------------------------
@pytest.mark.parametrize("text,expected", [
    ("I want to be a data analyst, contact me anytime", []),                  # Week 3 gave B.E. + M.E.
    ("Proficient in MS Office and Excel", []),                                # Week 3 gave M.S. in Office
    ("B.Tech Data Science with 3 years experience", ["B.Tech in Data Science"]),
    ("btech in computer science", ["B.Tech in Computer Science"]),
    ("M.S. in Computer Science", ["M.S. in Computer Science"]),
    ("BE in Electronics", ["B.E. in Electronics"]),
    ("MBA (Finance)", ["MBA in Finance"]),
    ("CFA level 2", ["CFA"]),
])
def test_degree_extraction(text, expected):
    assert [d.text for d in extract_degrees(text)] == expected


def test_phone_rejects_year_runs():
    assert extract_phones("worked 2019 2020 2021 and 2015-2016-2017") == []
    got = extract_phones("call +91 98765 43210 or 9876543210")
    assert len(got) == 2 and got[0].confidence > got[1].confidence - 0.2


def test_extract_response(nlp):
    r = nlp.extract("Jane, a@b.com, +91 98765 43210. B.Tech in Data Science, python, sql. Python again. 3 years experience.")
    names = {s.name: s for s in r.skills}
    assert names["python"].mentions == 2 and names["python"].confidence >= 0.9
    assert r.emails[0].value == "a@b.com" and r.experience_years == 3.0
    assert r.entities_count == len(r.skills) + len(r.degrees_or_certifications) + len(r.emails) + len(r.phone_numbers) + len(r.experience_mentions)




def test_academic_record_heading_preserves_first_subject_grade():
    text = (
        "Demo Student\n"
        "Degree: B.Tech in Data Science.\n"
        "Skills: Python, SQL, pandas, Power BI.\n"
        "Academic record: Physics: 42%; Mathematics: 76%; Chemistry: 68%; Overall GPA: 3.2."
    )
    records = extract_academic_records(text)
    by_subject = {record.subject: record for record in records}

    assert by_subject["Physics"].grade == "42%"
    assert by_subject["Physics"].score_percent == 42.0
    assert by_subject["Physics"].is_poor is True
    assert by_subject["Mathematics"].is_poor is False
    assert by_subject["Chemistry"].is_poor is False
    assert by_subject["Overall GPA"].score_percent == 80.0


def test_academic_grade_extraction_and_poor_grade_rules():
    records = extract_academic_records(
        "B.Tech in Data Science\nPhysics: 35%\nMathematics: 82%\nChemistry: D"
    )
    by_subject = {record.subject: record for record in records}
    assert by_subject["Physics"].score_percent == 35.0
    assert by_subject["Physics"].is_poor is True
    assert by_subject["Mathematics"].is_poor is False
    assert by_subject["Chemistry"].is_poor is True

    inline = extract_academic_records("I scored 38% in Physics and 75% in English.")
    assert [(r.subject, r.is_poor) for r in inline] == [("Physics", True), ("English", False)]
    grade_first = extract_academic_records("Grade: D in Physics")
    assert len(grade_first) == 1 and grade_first[0].subject == "Physics" and grade_first[0].is_poor
    inline_subjects = extract_academic_records("Physics: 35%, Mathematics: 80%")
    assert [(r.subject, r.is_poor) for r in inline_subjects] == [("Physics", True), ("Mathematics", False)]
    explicit = extract_academic_records("poor grades in Physics")
    assert len(explicit) == 1 and explicit[0].subject == "Physics" and explicit[0].is_poor
    assert extract_academic_records("Python, SQL, and 2 years of experience") == []


def test_academic_grades_support_percent_gpa_scales_and_letter_grades():
    records = extract_academic_records(
        "Physics: 42%\nMathematics GPA: 1.5/4\nChemistry GPA: 7/10\n"
        "Statistics GPA: 1.6\nBiology: Grade D\nOverall GPA: 3.2"
    )
    by_subject = {record.subject: record for record in records}

    assert by_subject["Physics"].score_percent == 42.0
    assert by_subject["Mathematics"].scale_max == 4.0
    assert by_subject["Mathematics"].score_percent == 37.5
    assert by_subject["Mathematics"].is_poor is True
    assert by_subject["Chemistry"].scale_max == 10.0
    assert by_subject["Chemistry"].score_percent == 70.0
    assert by_subject["Statistics"].scale_max == 4.0
    assert "assumed a 4-point GPA scale" in by_subject["Statistics"].evidence
    assert by_subject["Statistics"].score_percent == 40.0
    assert by_subject["Biology"].grade == "D"
    assert by_subject["Biology"].is_poor is True
    assert by_subject["Overall GPA"].scope == "overall"
    assert by_subject["Overall GPA"].scale_max == 4.0
    assert by_subject["Overall GPA"].score_percent == 80.0

    # An unscaled GPA defaults to /4 through 4. Values above 4 are inferred
    # to use the /10 scale, so 8.6 becomes 86% rather than an invalid /4 GPA.
    inferred_ten = extract_academic_records("GPA: 8.6")[0]
    assert inferred_ten.scale_max == 10.0
    assert inferred_ten.score_percent == 86.0
    assert inferred_ten.is_poor is False
    assert "assumed a 10-point GPA scale because the value exceeds 4" in inferred_ten.evidence


@pytest.mark.parametrize(
    "text,subject,expected_grade,expected_percent,expected_scale,expected_poor",
    [
        ("Physics = 42 percent", "Physics", "42%", 42.0, None, True),
        ("Physics 42 percent", "Physics", "42%", 42.0, None, True),
        ("Physics score 45 out of 100", "Physics", "45/100", 45.0, 100.0, True),
        ("Physics: GPA 1.5/4", "Physics", "GPA 1.5/4", 37.5, 4.0, True),
        ("Physics GPA: 7.0 out of 10", "Physics", "GPA 7/10", 70.0, 10.0, False),
        ("GPA: 3.0 on a 4-point scale", "Overall GPA", "GPA 3/4", 75.0, 4.0, False),
        ("CGPA 7.5 on a 10-point scale", "Overall GPA", "GPA 7.5/10", 75.0, 10.0, False),
        ("Grade A- in Physics", "Physics", "A-", None, None, False),
        ("Physics grade: B+", "Physics", "B+", None, None, False),
        ("Physics: F", "Physics", "F", None, None, True),
    ],
)
def test_academic_extraction_supports_common_cv_grade_formats(
    text, subject, expected_grade, expected_percent, expected_scale, expected_poor
):
    records = extract_academic_records(text)
    assert len(records) == 1
    record = records[0]
    assert record.subject == subject
    assert record.grade == expected_grade
    assert record.score_percent == expected_percent
    assert record.scale_max == expected_scale
    assert record.is_poor is expected_poor


def test_subject_gpa_without_scale_defaults_to_four_point_scale():
    records = extract_academic_records("Physics GPA: 1.5")
    assert len(records) == 1
    record = records[0]
    assert record.subject == "Physics"
    assert record.scale_max == 4.0
    assert record.score_percent == 37.5
    assert record.is_poor is True
    assert "assumed a 4-point GPA scale" in record.evidence




@pytest.mark.parametrize(
    "text,expected_scale,expected_percent,expected_poor,expected_evidence",
    [
        ("GPA: 3.7", 4.0, 92.5, False, "assumed a 4-point GPA scale"),
        ("GPA: 4.0", 4.0, 100.0, False, "assumed a 4-point GPA scale"),
        ("GPA: 4.1", 10.0, 41.0, True, "assumed a 10-point GPA scale because the value exceeds 4"),
        ("CGPA: 7.8", 10.0, 78.0, False, "assumed a 10-point GPA scale because the value exceeds 4"),
        ("GPA: 10", 10.0, 100.0, False, "assumed a 10-point GPA scale because the value exceeds 4"),
    ],
)
def test_unscaled_gpa_infers_four_or_ten_point_scale_from_value(
    text, expected_scale, expected_percent, expected_poor, expected_evidence
):
    records = extract_academic_records(text)
    assert len(records) == 1
    record = records[0]
    assert record.scale_max == expected_scale
    assert record.score_percent == expected_percent
    assert record.is_poor is expected_poor
    assert expected_evidence in record.evidence


def test_unscaled_gpa_above_ten_is_preserved_but_not_scored():
    record = extract_academic_records("GPA: 11.2")[0]
    assert record.scale_max == 10.0
    assert record.score_percent is None
    assert record.is_poor is False
    assert "value outside the scale; not normalised" in record.evidence


def test_subject_gpa_above_four_without_scale_uses_ten_point_scale():
    records = extract_academic_records("Physics GPA: 7.8")
    assert len(records) == 1
    record = records[0]
    assert record.subject == "Physics"
    assert record.grade == "GPA 7.8/10"
    assert record.scale_max == 10.0
    assert record.score_percent == 78.0
    assert record.is_poor is False
    assert "assumed a 10-point GPA scale because the value exceeds 4" in record.evidence


def test_explicit_gpa_scale_takes_precedence_over_inference():
    record = extract_academic_records("GPA: 8.6/4")[0]
    assert record.scale_max == 4.0
    assert record.score_percent is None
    assert record.is_poor is False
    assert "value outside the scale; not normalised" in record.evidence


def test_academic_parser_preserves_invalid_numeric_grades_without_scoring_them():
    records = extract_academic_records("Physics: 135%\nMaths: 12/10")
    by_subject = {record.subject: record for record in records}
    assert by_subject["Physics"].grade == "135%"
    assert by_subject["Physics"].score_percent is None
    assert by_subject["Physics"].is_poor is False
    assert by_subject["Maths"].grade == "12/10"
    assert by_subject["Maths"].score_percent is None
    assert by_subject["Maths"].scale_max == 10.0
    assert by_subject["Maths"].is_poor is False


def test_academic_parser_does_not_read_subject_letters_as_grades():
    records = extract_academic_records("Poor grade in Physics")
    assert len(records) == 1
    assert records[0].subject == "Physics"
    assert records[0].is_poor is True


def test_academic_extraction_uses_configured_poor_threshold(jobs):
    n = CareerNLP(
        AssetConfig(embedding_backend="lsa", lsa_dimensions=16, poor_grade_threshold_percent=60.0),
        jobs=jobs,
    ).warmup()
    result = n.extract("Physics: 55%")
    assert result.academic_records[0].is_poor is True


def test_semantic_requirement_fallback_handles_varied_jd_wording(nlp, monkeypatch):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 35%."
    called = []

    def semantic_fallback(st, subject, clause):
        called.append((subject, clause))
        return subject == "Physics" and "centerpiece" in clause.lower()

    monkeypatch.setattr(nlp, "_semantic_requirement_supported", semantic_fallback)
    result = nlp.match(
        student,
        "Data analyst. Physics is the centerpiece of this position. Python and SQL.",
    )
    assert result.academic_penalty_percent == 60.0
    assert "domain relevance 100%" in result.academic_adjustments[0].reason
    assert "semantic comparison with strong-requirement examples" in result.academic_adjustments[0].reason
    assert called and called[0][0] == "Physics"


@pytest.mark.parametrize(
    "job_clause,expected_penalty",
    [
        ("Candidates must be adept at Physics.", 40.0),
        ("We are looking for an expert in Physics.", 40.0),
        ("Physics specialist is the role we need.", 60.0),
        ("The candidate needs strong working knowledge of Physics.", 40.0),
        ("The role is built around applying Physics to complex problems.", 60.0),
        ("Advanced work in Physics is the main responsibility.", 60.0),
        ("Physics is not the main responsibility.", 0.0),
        ("The position is not centered on Physics.", 0.0),
        ("Physics experience is preferred but not mandatory.", 0.0),
        ("A Physics expert would be preferred, but it is not essential.", 0.0),
        ("Physics is mentioned; basic familiarity is sufficient.", 0.0),
    ],
)
def test_requirement_interpretation_covers_paraphrases_and_exclusions(
    nlp, job_clause, expected_penalty
):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 35%."
    result = nlp.match(student, f"Data analyst. {job_clause} Python and SQL.")
    assert result.academic_penalty_percent == expected_penalty


def test_mandatory_expertise_and_academic_foundation_triggers_subject_penalty(nlp):
    student = (
        "Demo Student\n"
        "Degree: B.Tech in Data Science.\n"
        "Skills: Python, SQL, pandas, Power BI, Excel.\n"
        "Academic record: Physics: 42%; Mathematics: 76%; Chemistry: 68%; Overall GPA: 3.2."
    )
    job_text = (
        "Work with physics simulations, mechanics and experimental physics datasets. "
        "Deep expertise in Physics and a strong academic foundation in Physics are mandatory. "
        "Python and statistical modelling are useful."
    )

    result = nlp.match(student, job_text, include_details=True)
    physics = next(record for record in result.academic_records if record.subject == "Physics")

    assert physics.is_poor is True
    assert result.academic_penalty_percent == 60.0
    assert [adjustment.subject for adjustment in result.academic_adjustments] == ["Physics"]
    assert result.academic_adjustments[0].domain_relevance_percent == 100.0
    assert result.match_score == pytest.approx(max(0.0, result.base_match_score - 60.0))


def test_match_penalises_only_a_strong_requirement_for_a_weak_subject(nlp):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 35%."
    requires_physics = (
        "Data analyst. Requires strong expertise in Physics. Skills: python, sql."
    )
    physics_not_required = (
        "Data analyst. Physics expertise is not required. Skills: python, sql."
    )

    risky = nlp.match(student, requires_physics, include_details=True)
    safe = nlp.match(student, physics_not_required, include_details=True)

    assert risky.academic_penalty_percent == 40.0
    assert risky.academic_adjustments[0].subject == "Physics"
    assert risky.match_score == pytest.approx(max(0, risky.base_match_score - 40.0))
    assert risky.score_breakdown.academic_penalty_percent == 40.0
    assert safe.academic_penalty_percent == 0.0
    assert safe.academic_adjustments == []


def test_core_domain_penalty_is_higher_than_non_core_subject_requirement(nlp):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 42%."
    results = nlp.compare_jobs(student, [
        {
            "job_id": "GENERAL",
            "job_title": "Data Analyst",
            "job_text": "Data analyst. Applicants must be adept at Physics. Skills: python, sql.",
        },
        {
            "job_id": "PHYSICS",
            "job_title": "Physics Research Data Analyst",
            "job_text": (
                "Work with physics simulations, mechanics and experimental physics datasets. "
                "Deep expertise in Physics is mandatory. Python and SQL are used for analysis."
            ),
        },
    ])
    ranked = {item.job_id: item.match for item in [results.recommended_job, *results.alternatives]}

    assert ranked["GENERAL"].academic_penalty_percent == 40.0
    assert ranked["PHYSICS"].academic_penalty_percent == 60.0
    assert ranked["GENERAL"].academic_adjustments[0].domain_relevance_percent == 0.0
    assert ranked["PHYSICS"].academic_adjustments[0].domain_relevance_percent == 100.0
    assert ranked["PHYSICS"].match_score == pytest.approx(
        max(0.0, ranked["PHYSICS"].base_match_score - 60.0)
    )


def test_core_domain_can_trigger_adjustment_without_keyword_mandatory_requirement(nlp):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 42%."
    response = nlp.compare_jobs(student, [
        {
            "job_id": "PHYSICS",
            "job_title": "Physics Research Data Analyst",
            "job_text": (
                "Work with physics simulations, mechanics and experimental physics datasets. "
                "Use Python and SQL to process research data."
            ),
        },
        {
            "job_id": "GENERAL",
            "job_title": "Business Intelligence Analyst",
            "job_text": "Build business dashboards using Power BI, SQL and Excel.",
        },
    ])
    physics = next(
        item.match for item in [response.recommended_job, *response.alternatives]
        if item.job_id == "PHYSICS"
    )

    assert physics.academic_penalty_percent == 60.0
    assert physics.academic_adjustments[0].domain_relevance_percent == 100.0
    assert "core-domain evidence" in physics.academic_adjustments[0].reason


def test_preferred_or_helpful_subject_expertise_does_not_trigger_penalty(nlp):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 35%."
    for job_text in [
        "Data analyst. Strong expertise in Physics is preferred. Python and SQL.",
        "Data analyst. Physics expertise would be helpful, but it is not required.",
    ]:
        result = nlp.match(student, job_text)
        assert result.academic_penalty_percent == 0.0


def test_academic_record_preserves_short_acronym_subject_names(nlp):
    records = nlp.extract("Academic record: NLP: 35%; RAG: 42%.").academic_records
    assert {record.subject for record in records} >= {"NLP", "RAG"}


def test_generic_competencies_detect_unknown_topics_and_preferred_knowledge(nlp):
    student = (
        "B.Tech in Data Science. Skills: Python, SQL, pandas. "
        "Projects: Built reporting dashboards and cleaned business datasets. "
        "Academic record: NLP: 35%; Chemistry: 35%."
    )
    response = nlp.compare_jobs(student, [
        {
            "job_id": "TITLE_MISMATCH_CORE_NLP",
            "job_title": "Business Intelligence Analyst",
            "job_text": (
                "Despite the title, the central technical domain of this role is "
                "Natural Language Processing (NLP). The work is built around RAG systems "
                "and NLP model evaluation. Deep expertise in NLP is required. "
                "Strong knowledge of RAG is mandatory. Candidates use Python, pandas, and SQL "
                "to process data. Knowledge of Chemistry is nice to have."
            ),
        },
        {
            "job_id": "TITLE_MISMATCH_NOT_CORE",
            "job_title": "NLP Research Engineer",
            "job_text": (
                "This is a business intelligence role focused on Power BI dashboards, SQL "
                "reporting, and sales trends. NLP expertise is not required. "
                "Knowledge of Chemistry is not required but good to have."
            ),
        },
    ])
    by_id = {
        item.job_id: item.match
        for item in [response.recommended_job, *response.alternatives]
    }
    core = by_id["TITLE_MISMATCH_CORE_NLP"]
    not_core = by_id["TITLE_MISMATCH_NOT_CORE"]
    requirements = {item.name.casefold(): item for item in core.competency_requirements}

    assert {"python", "sql", "pandas"} <= set(name.casefold() for name in core.matched_skills)
    assert {"nlp", "rag"} <= set(name.casefold() for name in core.missing_weak_skills)
    assert "chemistry" not in {name.casefold() for name in core.missing_weak_skills}
    assert {"nlp", "rag", "chemistry"} <= set(name.casefold() for name in core.recommended_learning)
    assert requirements["nlp"].priority == "required"
    assert requirements["nlp"].status == "missing"
    assert requirements["rag"].priority == "required"
    assert requirements["rag"].status == "missing"
    assert requirements["chemistry"].priority == "preferred"
    assert requirements["chemistry"].status == "missing"
    assert core.academic_penalty_percent == 60.0
    assert not_core.academic_penalty_percent == 0.0
    assert not any(
        item.name.casefold() == "nlp" and item.priority != "not_required"
        for item in not_core.competency_requirements
    )
    chemistry_preference = next(
        item for item in not_core.competency_requirements
        if item.name.casefold() == "chemistry"
    )
    assert chemistry_preference.priority == "preferred"
    assert chemistry_preference.status == "missing"



def test_preferred_competency_is_not_a_required_gap_but_is_learning_option(nlp):
    student = "B.Tech in Data Science. Skills: Python, SQL. Chemistry: 35%."
    result = nlp.match(
        student,
        "Data analyst requiring Python and SQL. Knowledge of Chemistry is nice to have."
    )
    chemistry = next(
        item for item in result.competency_requirements if item.name.casefold() == "chemistry"
    )
    assert chemistry.priority == "preferred"
    assert chemistry.status == "missing"
    assert "chemistry" not in {name.casefold() for name in result.missing_weak_skills}
    assert "chemistry" in {name.casefold() for name in result.recommended_learning}


def test_required_skill_count_excludes_preferred_and_not_required(nlp):
    result = nlp.match(
        "B.Tech in Data Science. Skills: Python, SQL.",
        "Python and SQL are required. Knowledge of Chemistry is nice to have. "
        "Knowledge of Biology is not required.",
        include_details=True,
    )
    assert result.score_breakdown.required_skills_detected == 2


def test_role_title_or_domain_description_does_not_create_skill_requirement(nlp):
    result = nlp.match(
        "B.Tech in Data Science. Skills: dashboards, reporting, SQL.",
        "This is a business intelligence role focused on Power BI dashboards, SQL reporting, and sales trends. "
        "The role does not require business intelligence expertise.",
    )
    requirements = {item.name.casefold(): item for item in result.competency_requirements}
    assert "business intelligence" not in requirements
    assert "power bi" in requirements
    assert "sql" in requirements


def test_explicit_domain_knowledge_requirement_is_still_extracted(nlp):
    result = nlp.match(
        "B.Tech in Data Science. Skills: Python, SQL.",
        "This is a business intelligence role. Strong knowledge of business intelligence concepts is required. "
        "Python and SQL are also required.",
    )
    requirement = next(
        item for item in result.competency_requirements
        if item.name.casefold() == "business intelligence"
    )
    assert requirement.priority == "required"
    assert requirement.status == "missing"
    assert "business intelligence" in {name.casefold() for name in result.missing_weak_skills}


def test_interest_only_or_intended_learning_is_not_competency_evidence(nlp):
    student = (
        "B.Tech in Data Science. I am interested in NLP and want to learn RAG. "
        "Academic record: NLP: 35%."
    )
    result = nlp.match(
        student,
        "Natural Language Processing (NLP) is the core technical domain of this role. "
        "Deep expertise in NLP is required. Strong knowledge of RAG is mandatory.",
    )
    requirements = {item.name.casefold(): item for item in result.competency_requirements}

    assert requirements["nlp"].status == "missing"
    assert requirements["rag"].status == "missing"
    assert {"nlp", "rag"} <= set(result.missing_weak_skills)
    assert {"nlp", "rag"} <= set(result.recommended_learning)
    assert result.academic_penalty_percent == 60.0


def test_demonstrated_subject_competence_suppresses_grade_penalty(nlp):
    student = (
        "B.Tech in Data Science. Skills: Python, SQL, Physics simulations. "
        "Project: Built and evaluated physics simulations with Python. Physics: 35%."
    )
    result = nlp.match(
        student,
        "Physics is the core technical domain of this role. "
        "Deep expertise in Physics is mandatory. Use Python to analyse simulation results.",
        include_details=True,
    )
    physics_requirement = next(
        item for item in result.competency_requirements
        if item.name.casefold() == "physics"
    )
    assert physics_requirement.status == "matched"
    assert "physics" in {name.casefold() for name in result.matched_skills}
    assert result.academic_penalty_percent == 0.0
    assert any(record.subject == "Physics" and record.is_poor for record in result.academic_records)


def test_compare_jobs_ranks_jd_that_does_not_require_weak_subject(nlp):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: 35%."
    response = nlp.compare_jobs(student, [
        {
            "job_id": "JD1",
            "job_title": "Physics-heavy analyst",
            "job_text": "Data analyst requiring strong expertise in Physics. Skills: python, sql.",
        },
        {
            "job_id": "JD2",
            "job_title": "General data analyst",
            "job_text": "Data analyst using python and sql. Physics expertise is not required.",
        },
    ])
    assert response.recommended_job.job_id == "JD2"
    assert response.alternatives[0].job_id == "JD1"
    assert response.recommended_job.match.academic_penalty_percent == 0.0
    assert response.alternatives[0].match.academic_penalty_percent == 40.0
    assert response.academic_records[0].subject == "Physics"


def test_exact_sample_case_reranks_away_from_mandatory_physics_role(tmp_path):
    """Exercise the reported CV/JDs against the full bundled job corpus offline."""
    student = (
        "Demo Student\n"
        "Degree: B.Tech in Data Science.\n"
        "Skills: Python, SQL, pandas, Power BI, Excel, data visualization, statistics, "
        "basic machine learning and data cleaning.\n"
        "Projects: Built a sales dashboard in Power BI, cleaned datasets with Python and SQL, "
        "and analysed business trends.\n"
        "Academic record: Physics: 42%; Mathematics: 76%; Chemistry: 68%; Overall GPA: 3.2."
    )
    jobs = [
        {
            "job_id": "JD1",
            "job_title": "Physics Research Data Analyst",
            "job_text": (
                "Work with physics simulations, mechanics and experimental physics datasets. "
                "Deep expertise in Physics and a strong academic foundation in Physics are mandatory. "
                "Python and statistical modelling are useful."
            ),
        },
        {
            "job_id": "JD2",
            "job_title": "Business Intelligence Analyst",
            "job_text": (
                "Build business dashboards and reporting pipelines using Power BI, SQL and Excel. "
                "Strong data visualisation, analytical thinking and stakeholder communication are required. "
                "Physics expertise is not required."
            ),
        },
        {
            "job_id": "JD3",
            "job_title": "Data Operations Analyst",
            "job_text": (
                "Clean and validate datasets, run SQL queries, use Python and pandas, and troubleshoot "
                "data-quality issues. Attention to detail and documentation skills are important. "
                "The role does not require Physics expertise."
            ),
        },
    ]
    engine = CareerNLP(
        AssetConfig(
            embedding_backend="lsa",
            cache_dir=tmp_path / "cache",
            lsa_dimensions=256,
        )
    ).warmup()

    response = engine.compare_jobs(student, jobs, include_details=True)
    ranked = [response.recommended_job, *response.alternatives]
    by_id = {item.job_id: item for item in ranked}
    physics = next(record for record in response.academic_records if record.subject == "Physics")

    assert physics.score_percent == 42.0 and physics.is_poor is True
    assert by_id["JD1"].match.academic_penalty_percent == 60.0
    assert by_id["JD2"].match.academic_penalty_percent == 0.0
    assert by_id["JD3"].match.academic_penalty_percent == 0.0
    assert response.recommended_job.job_id in {"JD2", "JD3"}
    assert by_id["JD1"].match.match_score == pytest.approx(
        max(0.0, by_id["JD1"].match.base_match_score - 60.0)
    )


def test_compare_jobs_validates_count_and_duplicate_ids(nlp):
    with pytest.raises(InputValidationError):
        nlp.compare_jobs("Python", [{"job_text": "First JD"}])
    with pytest.raises(InputValidationError):
        nlp.compare_jobs("Python", [
            {"job_id": "same", "job_text": "First JD"},
            {"job_id": "same", "job_text": "Second JD"},
        ])

def test_skill_confidence_rules():
    assert skill_confidence("r", 1) < skill_confidence("sap", 1) < skill_confidence("power bi", 1)
    assert skill_confidence("python", 2) > skill_confidence("python", 1)


def test_text_helpers():
    assert clean_job_title("data solutions analyst ii (managed care - finance focus) - r7269") == "data solutions analyst"
    assert parse_skills("['Power  BI', 'SQL', 'sql']") == ["power bi", "sql"]
    assert parse_skills("a; b") == ["a", "b"]


# --------------------------------------------------------------------------
# validation + errors
# --------------------------------------------------------------------------
@pytest.mark.parametrize("bad", ["", "   ", None, 123])
def test_rejects_bad_text(nlp, bad):
    with pytest.raises(InputValidationError):
        nlp.extract(bad)


def test_rejects_oversized_and_bad_topk(jobs):
    small = CareerNLP(AssetConfig(embedding_backend="lsa", lsa_dimensions=16, max_text_chars=20), jobs=jobs)
    with pytest.raises(InputValidationError) as ei:
        small.extract("x" * 21)
    assert ei.value.details["max"] == 20
    with pytest.raises(InputValidationError):
        small.recommend("python", top_k=0)


def test_bad_dataset_errors(tmp_path):
    with pytest.raises(ModelLoadError):
        CareerNLP(AssetConfig(jobs_path=tmp_path / "nope.csv", embedding_backend="lsa")).warmup()
    with pytest.raises(ModelLoadError) as ei:
        CareerNLP(AssetConfig(embedding_backend="lsa"), jobs=pd.DataFrame({"job_id": ["a"]})).warmup()
    assert "category" in ei.value.details["missing"]


def test_config_validation():
    with pytest.raises(ValueError):
        AssetConfig(skill_weight=0.5, semantic_weight=0.6)
    with pytest.raises(ValueError):
        AssetConfig(embedding_backend="bogus")
    assert AssetConfig(academic_subject_penalty=70.0).academic_core_domain_penalty == 70.0
    with pytest.raises(ValueError):
        AssetConfig(academic_core_domain_penalty=101.0)


def test_config_from_env(monkeypatch, tmp_path):
    monkeypatch.setenv("CAREER_NLP_JOBS_PATH", str(tmp_path / "j.csv"))
    monkeypatch.setenv("CAREER_NLP_BACKEND", "lsa")
    monkeypatch.setenv("CAREER_NLP_WARMUP", "false")
    monkeypatch.setenv("CAREER_NLP_ACADEMIC_CORE_DOMAIN_PENALTY", "65")
    c = AssetConfig.from_env()
    assert c.jobs_path == tmp_path / "j.csv"
    assert c.embedding_backend == "lsa"
    assert c.warmup_on_start is False
    assert c.academic_core_domain_penalty == 65.0


# --------------------------------------------------------------------------
# pluggable embedder, cache, Sentence-Transformer loader (stubbed)
# --------------------------------------------------------------------------
def test_custom_embedder_and_disk_cache(jobs, tmp_path):
    cfg = AssetConfig(cache_dir=tmp_path / "cache")
    a = CareerNLP(cfg, jobs=jobs, embedder=HashEmbedder()).warmup()
    files = list((tmp_path / "cache").glob("job_index_*.npz"))
    assert len(files) == 1
    mtime = files[0].stat().st_mtime_ns
    b = CareerNLP(cfg, jobs=jobs, embedder=HashEmbedder()).warmup()   # reuses cache
    assert files[0].stat().st_mtime_ns == mtime
    assert np.allclose(a._state.index_vectors, b._state.index_vectors)
    assert a.match(IT_STUDENT, IT_JD).meta.model.embedding_backend == "custom"
    # changing the data invalidates the cache key
    CareerNLP(cfg, jobs=jobs.iloc[:-1], embedder=HashEmbedder()).warmup()
    assert len(list((tmp_path / "cache").glob("job_index_*.npz"))) == 2


def test_auto_backend_falls_back_with_warning(jobs, monkeypatch):
    monkeypatch.setitem(sys.modules, "sentence_transformers", None)   # force ImportError
    n = CareerNLP(AssetConfig(embedding_backend="auto", lsa_dimensions=16), jobs=jobs)
    r = n.extract("python")
    assert r.meta.model.embedding_backend == "lsa"
    assert any("LSA offline fallback" in w for w in r.meta.warnings)
    strict = CareerNLP(AssetConfig(embedding_backend="sentence-transformers"), jobs=jobs)
    with pytest.raises(ModelLoadError):
        strict.warmup()


def test_sentence_transformer_loader_uses_local_snapshot(tmp_path, monkeypatch):
    snap = tmp_path / "models--sentence-transformers--all-MiniLM-L6-v2" / "snapshots" / "abc"
    snap.mkdir(parents=True)
    (snap / "config.json").write_text("{}"); (snap / "model.safetensors").write_text("x")
    assert find_local_snapshot("all-MiniLM-L6-v2", tmp_path) == snap
    assert find_local_snapshot("other-model", tmp_path) is None

    calls = []

    class FakeST:
        def __init__(self, path, **kw):
            calls.append((path, kw))
        def get_sentence_embedding_dimension(self):
            return 384
        def encode(self, texts, **kw):
            return np.ones((len(texts), 384), dtype=np.float32) / np.sqrt(384)

    monkeypatch.setitem(sys.modules, "sentence_transformers", types.SimpleNamespace(SentenceTransformer=FakeST))
    emb = SentenceTransformerEmbedder("all-MiniLM-L6-v2", tmp_path)
    assert calls == [(str(snap), {})]                 # loaded from disk, no download args
    assert emb.dimensions == 384 and emb.encode(["a", "b"]).shape == (2, 384)
    # no snapshot -> hub name + cache_folder (download path)
    SentenceTransformerEmbedder("all-MiniLM-L6-v2", tmp_path / "empty")
    assert calls[-1][0] == "all-MiniLM-L6-v2" and "cache_folder" in calls[-1][1]


# --------------------------------------------------------------------------
# CV
# --------------------------------------------------------------------------
def _docx_bytes(paragraphs, table_cells=()):
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table_cells:
        t = d.add_table(rows=1, cols=len(table_cells))
        for c, txt in zip(t.rows[0].cells, table_cells):
            c.text = txt
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def test_cv_docx_including_tables(nlp):
    data = _docx_bytes(
        ["Jane Doe  jane@x.com", "B.Tech in Data Science. 2 years experience.", "Physics: 42%."],
        table_cells=["Skills", "python, sql, machine learning"],
    )
    r = nlp.analyze_cv_bytes(data, "jane.docx")
    assert r.filename == "jane.docx" and {"python", "sql"} <= {s.name for s in r.entities.skills}
    assert r.profile.normalized_education == "B.Tech in Data Science"
    assert r.entities.academic_records[0].subject == "Physics"
    assert r.entities.academic_records[0].is_poor is True


def test_cv_errors(nlp):
    with pytest.raises(CVProcessingError):
        nlp.analyze_cv_bytes(b"hello", "cv.txt")
    with pytest.raises(CVProcessingError):
        nlp.analyze_cv_bytes(b"not really a pdf", "cv.pdf")
    with pytest.raises(CVProcessingError):
        nlp.analyze_cv_bytes(_docx_bytes([""]), "empty.docx")
    with pytest.raises(CVProcessingError):
        nlp.analyze_cv("/no/such/file.pdf")


# --------------------------------------------------------------------------
# REST API
# --------------------------------------------------------------------------
@pytest.fixture(scope="module")
def client(nlp):
    from fastapi.testclient import TestClient
    from career_nlp.api import create_app
    return TestClient(create_app(engine=nlp))


def test_api_health_info(client):
    h = client.get("/v1/health").json()
    assert h["status"] == "ok" and h["ready"] and h["jobs_indexed"] == 48
    i = client.get("/v1/info").json()
    assert i["scoring"] == {
        "skill_weight": 0.63,
        "semantic_weight": 0.27,
        "category_weight": 0.1,
        "poor_grade_threshold_percent": 50.0,
        "academic_subject_penalty": 40.0,
        "academic_core_domain_penalty": 60.0,
        "preferred_requirement_weight": 0.35,
    } and "HR" in i["categories"]


def test_api_match_and_recommend(client):
    r = client.post("/v1/match", json={"student_text": IT_STUDENT, "job_text": IT_JD})
    assert r.status_code == 200 and r.headers["X-Request-ID"]
    body = r.json()
    assert "score_breakdown" not in body and body["confidence"]["level"]
    d = client.post("/v1/match", json={"student_text": IT_STUDENT, "job_text": IT_JD, "include_details": True}).json()
    assert d["score_breakdown"]["skill_weight"] == 0.63
    rec = client.post("/v1/recommend", json={"student_text": IT_STUDENT, "top_k": 2}).json()
    assert rec["recommended_job"]["category"] == "INFORMATION-TECHNOLOGY" and len(rec["alternatives"]) == 1
    both = client.post("/v1/analyze", json={"student_text": IT_STUDENT}).json()
    assert both["flow"] == "jd_not_provided" and "match" not in both


@pytest.mark.parametrize("path,payload", [
    ("/v1/extract", {"text": IT_STUDENT}), ("/v1/classify", {"text": IT_STUDENT, "top_k": 2}),
    ("/v1/embed", {"texts": ["a", "b"]}), ("/v1/similarity", {"text_a": "a b", "text_b": "a b"}),
    ("/v1/summarize", {"text": "One. Two. Three. Four."}), ("/v1/retrieve", {"query": "python", "top_k": 2}),
])
def test_api_nlp_endpoints(client, path, payload):
    r = client.post(path, json=payload)
    assert r.status_code == 200, r.text
    assert r.json()["meta"]["schema_version"] == "1.0"


def test_api_match_and_analyze_apply_subject_grade_adjustment(client):
    student = "B.Tech in Data Science. Skills: python, sql. Physics: GPA 1.2/4."
    jd = "Data analyst. Applicants must be adept at Physics. Skills: python, sql."

    direct_match = client.post("/v1/match", json={"student_text": student, "job_text": jd})
    assert direct_match.status_code == 200, direct_match.text
    match_payload = direct_match.json()
    assert match_payload["academic_records"][0]["subject"] == "Physics"
    assert match_payload["academic_records"][0]["scale_max"] == 4.0
    assert match_payload["academic_penalty_percent"] == 40.0

    analyze = client.post("/v1/analyze", json={"student_text": student, "job_text": jd})
    assert analyze.status_code == 200, analyze.text
    assert analyze.json()["flow"] == "jd_provided"
    assert analyze.json()["match"]["academic_penalty_percent"] == 40.0


def test_api_compare_jobs_and_grade_extraction(client):
    student = "B.Tech in Data Science. Skills: python, sql. Academic record: Physics: 35%."
    extracted = client.post("/v1/extract", json={"text": student})
    assert extracted.status_code == 200
    assert extracted.json()["academic_records"][0]["subject"] == "Physics"
    assert extracted.json()["academic_records"][0]["is_poor"] is True
    recommendation = client.post("/v1/recommend", json={"student_text": student, "top_k": 2})
    assert recommendation.status_code == 200
    assert recommendation.json()["academic_records"][0]["subject"] == "Physics"

    response = client.post("/v1/compare-jobs", json={
        "student_text": student,
        "job_descriptions": [
            {"job_id": "JD1", "job_title": "Physics-heavy analyst",
             "job_text": "Data analyst requiring strong expertise in Physics. Skills: python, sql."},
            {"job_id": "JD2", "job_title": "General data analyst",
             "job_text": "Data analyst using python and sql. Physics expertise is not required."},
        ],
    })
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["recommended_job"]["job_id"] == "JD2"
    assert payload["alternatives"][0]["match"]["academic_penalty_percent"] == 40.0

    invalid = client.post("/v1/compare-jobs", json={
        "student_text": student,
        "job_descriptions": [{"job_text": "Only one JD"}],
    })
    assert invalid.status_code == 422


def test_api_error_contract(client):
    r = client.post("/v1/match", json={"student_text": "", "job_text": "x"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_request"
    assert r.json()["error"]["request_id"] == r.headers["X-Request-ID"]
    r = client.post("/v1/match", json={"student_text": "   ", "job_text": "x"})   # passes pydantic, fails engine
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_input"
    r = client.post("/v1/match", json={"student_text": "a", "job_text": "b", "bogus": 1})
    assert r.status_code == 422                                                   # extra fields forbidden
    r = client.post("/v1/recommend", json={"student_text": "a", "top_k": 99})
    assert r.status_code == 422
    assert client.get("/v1/nope").status_code == 404


def test_api_cv_upload(client):
    data = _docx_bytes(["B.Tech in Data Science", "python sql", "Physics GPA: 1.5/4"])
    r = client.post("/v1/cv/analyze", files={"file": ("cv.docx", data, "application/octet-stream")})
    assert r.status_code == 200 and r.json()["profile"]["skills"]
    assert r.json()["entities"]["academic_records"][0]["subject"] == "Physics"
    assert r.json()["entities"]["academic_records"][0]["is_poor"] is True
    r = client.post("/v1/cv/analyze", files={"file": ("cv.exe", b"x", "application/octet-stream")})
    assert r.status_code == 422 and r.json()["error"]["code"] == "cv_processing_failed"


def test_api_key_enforced(nlp, monkeypatch):
    from fastapi.testclient import TestClient
    from career_nlp.api import create_app
    monkeypatch.setenv("CAREER_NLP_API_KEY", "s3cret")
    c = TestClient(create_app(engine=nlp))
    assert c.get("/v1/health").status_code == 200                                  # health stays open
    assert c.get("/v1/info").status_code == 401
    assert c.get("/v1/info", headers={"X-API-Key": "wrong"}).status_code == 401
    assert c.get("/v1/info", headers={"X-API-Key": "s3cret"}).status_code == 200


def test_api_health_degraded_before_warmup(jobs):
    from fastapi.testclient import TestClient
    from career_nlp.api import create_app
    cold = CareerNLP(AssetConfig(embedding_backend="lsa", lsa_dimensions=16, warmup_on_start=False), jobs=jobs)
    c = TestClient(create_app(engine=cold))
    h = c.get("/v1/health").json()
    assert h["status"] == "degraded" and h["ready"] is False
    assert c.post("/v1/extract", json={"text": "python"}).status_code == 200      # lazy-inits on demand
    assert c.get("/v1/health").json()["ready"] is True


def test_cli_roundtrip(jobs, tmp_path, capsys):
    from career_nlp.cli import main
    p = tmp_path / "jobs.csv"; jobs.to_csv(p, index=False)
    assert main(["--backend", "lsa", "--jobs", str(p), "extract", "--text", "python and sql"]) == 0
    assert "python" in json.loads(capsys.readouterr().out)["skills"][0]["name"]
    assert main(["--backend", "lsa", "--jobs", str(p), "match", "--student", "", "--job", "x"]) == 2
    assert json.loads(capsys.readouterr().err)["error"]["code"] == "invalid_input"
